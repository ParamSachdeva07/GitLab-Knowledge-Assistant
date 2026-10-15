from pathlib import Path
from langchain_core.documents import Document
from langchain_text_splitters import MarkdownTextSplitter, RecursiveCharacterTextSplitter
from src.chunking.loader import normalize_metadata, add_chunks
from src.ingestion.filter import filter_files
from src.ingestion.sync import sync_data, save_json
import json
import fcntl
from datetime import datetime, timezone
from uuid import uuid4
from langfuse import observe, get_client
import csv
from io import StringIO
from openpyxl import load_workbook
import logging
from frontmatter.default_handlers import YAMLHandler
from yaml import YAMLError
from docling.document_converter import DocumentConverter
from docling.datamodel.base_models import ConversionStatus
from src.db.client import client, collection
from qdrant_client import models

project_root = Path(__file__).resolve().parents[2]
repo = project_root / "data" / "handbook"
logger = logging.getLogger(__name__)


class PDFConversionError(RuntimeError):
    pass

def load_documents_md(paths):
    return [
        Document(
            page_content=path.read_text(encoding='utf-8'),
            metadata={'source': str(path)}
        )
        for path in paths
    ]

def load_documents_yaml(paths):
    return [
        Document(
            page_content=path.read_text(encoding='utf-8'),
            metadata={
                'source': str(path),
                "format": "yaml"
            }
        ) for path in paths
    ]

@observe(name="ingestion.pdf_conversion", capture_input=False, capture_output=False)
def convert_pdf(path):
    get_client().update_current_span(metadata={"source": str(path)})
    try:
        result = DocumentConverter().convert(path)
        if result.status != ConversionStatus.SUCCESS:
            raise RuntimeError(f"Conversion status: {result.status}")

        documents = []
        for page_number in sorted(result.document.pages):
            text = result.document.export_to_markdown(page_no=page_number)
            if text.strip():
                documents.append(Document(
                    page_content=text,
                    metadata={"source": str(path), "format": "pdf", "page_number": page_number},
                ))
        if not documents:
            raise RuntimeError("No text extracted from PDF")
        get_client().update_current_span(output={"pages_with_text": len(documents)})
        return documents
    except Exception as error:
        raise PDFConversionError(f"PDF conversion failed for {path}: {error}") from error


def load_documents_pdf(paths):
    documents = []
    for path in paths:
        documents.extend(convert_pdf(path))
    return documents


def extract_frontmatter(documents):
    handler = YAMLHandler()
    for document in documents:
        text = document.page_content.strip()
        if not handler.detect(text):
            continue
        try:
            raw_metadata, content = handler.split(text)
            yaml_metadata = handler.load(raw_metadata)
        except (YAMLError, ValueError) as error:
            logger.warning("Ignoring invalid frontmatter in %s: %s", document.metadata.get("source"), type(error).__name__)
            document.metadata["frontmatter_status"] = "invalid"
            continue

        if not isinstance(yaml_metadata, dict):
            logger.warning("Ignoring non-mapping frontmatter in %s", document.metadata.get("source"))
            document.metadata["frontmatter_status"] = "unsupported_type"
            continue

        # Reject non-string top-level keys; nested JSON-compatible values remain valid.
        yaml_metadata = {key: value for key, value in yaml_metadata.items() if isinstance(key, str)}
        document.page_content = content.strip()
        document.metadata = {
            **document.metadata,
            **normalize_metadata(yaml_metadata),
            "source": document.metadata["source"],
        }
    return documents


def chunk_markdown_documents(documents):
    splitter = MarkdownTextSplitter()

    chunks = splitter.split_documents(documents)
    return chunks

def load_and_chunk_xlsx(paths, max_rows=10, max_chars=4000):
    """Group whole rows, repeating row-1 headers. A single row may exceed max_chars."""
    if max_rows < 1 or max_chars < 1:
        raise ValueError("max_rows and max_chars must be positive")

    def row_to_text(values):
        buffer = StringIO()
        csv.writer(buffer, lineterminator="\n").writerow(values)
        return buffer.getvalue()

    chunks = []

    for path in paths:
        path = Path(path)
        if path.name.startswith("~$"):
            continue

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook:
                rows = sheet.iter_rows(values_only=True)
                headers = list(next(rows, ()))

                while headers and headers[-1] is None:
                    headers.pop()
                if not headers:
                    continue

                prefix = f"Sheet: {sheet.title}\n" + row_to_text(headers)
                text = prefix
                row_numbers = []

                def save_chunk():
                    chunks.append(
                        Document(
                            page_content=text,
                            metadata={
                                "source": str(path),
                                "format": "xlsx",
                                "sheet": sheet.title,
                                "row_start": row_numbers[0],
                                "row_end": row_numbers[-1],
                            },
                        )
                    )

                for row_number, row in enumerate(rows, start=2):
                    if all(value is None for value in row):
                        continue

                    if any(value is not None for value in row[len(headers):]):
                        raise ValueError(
                            f"Data beyond headers: {path}, {sheet.title}, row {row_number}"
                        )

                    row_text = row_to_text(row[:len(headers)])
                    would_exceed = (
                        len(row_numbers) >= max_rows
                        or len(text) + len(row_text) > max_chars
                    )
                    if row_numbers and would_exceed:
                        save_chunk()
                        text = prefix
                        row_numbers = []

                    text += row_text
                    row_numbers.append(row_number)

                if row_numbers:
                    save_chunk()
        finally:
            workbook.close()

    return chunks


def prepare_file(path):
    if path.suffix.lower() == ".md":
        chunks = chunk_markdown_documents(extract_frontmatter(load_documents_md([path])))
    elif path.suffix.lower() == ".pdf":
        chunks = chunk_markdown_documents(load_documents_pdf([path]))
    elif path.suffix.lower() in {".yaml", ".yml"}:
        chunks = RecursiveCharacterTextSplitter().split_documents(load_documents_yaml([path]))
    else:
        chunks = load_and_chunk_xlsx([path])
    return chunks


def source_filter_for(paths):
    sources = set()
    for path in paths:
        path = Path(path)
        relative_path = path.relative_to(repo)
        sources.update({
            str(path),
            str(Path("../../data/handbook") / relative_path),
            str(Path("../../data/cleaned") / relative_path),
            str(project_root / "data" / "cleaned" / relative_path),
        })
    return models.Filter(must=[models.FieldCondition(
        key="metadata.source", match=models.MatchAny(any=sorted(sources)),
    )])


def finish_pending_file(pending_path, queue, path, result):
    if pending_path is None:
        return
    pending = json.loads(pending_path.read_text())
    pending[queue].remove(str(path))
    pending["completed_files"][f"{queue}:{path}"] = result
    save_json(pending_path, pending)


def delete_sources(paths, pending_path=None):
    """Delete each source and acknowledge it only after verification."""
    deleted = 0
    for path in paths:
        path = Path(path)
        try:
            source_filter = source_filter_for([path])
            count = client.count(collection, count_filter=source_filter, exact=True).count
            client.delete(
                collection_name=collection,
                points_selector=models.FilterSelector(filter=source_filter),
                wait=True,
            )
            if client.count(collection, count_filter=source_filter, exact=True).count:
                raise RuntimeError("Chunks remain after deletion")
            deleted += count
            finish_pending_file(pending_path, "remaining_deletions", path, {"chunks_deleted": count})
        except Exception as error:
            raise RuntimeError(f"Deletion failed for {path}: {error}") from error
    return deleted


@observe(name="ingestion.add_sources", capture_input=False, capture_output=False)
def add_sources(paths, pending_path=None):
    """Prepare all pending files; acknowledge them only after insertion succeeds."""
    paths = [Path(path) for path in paths]
    all_chunks = []
    results = {}

    for path in paths:
        try:
            chunks = prepare_file(path)
        except PDFConversionError as error:
            logger.warning("Skipping PDF: %s", error)
            results[path] = {"status": "skipped", "reason": "pdf_conversion_failed", "error": str(error)}
            continue
        except Exception as error:
            raise RuntimeError(f"Preparation failed for {path}: {error}") from error
        all_chunks.extend(chunks)
        results[path] = {"status": "indexed", "chunks": len(chunks)}

    # Clear partial writes from an earlier attempt without acknowledging queues.
    delete_sources(paths)
    usage = add_chunks(all_chunks) if all_chunks else {"input_tokens": 0}

    # Verify every source before acknowledging any successful additions.
    for path, result in results.items():
        if result["status"] == "indexed":
            actual = client.count(collection, count_filter=source_filter_for([path]), exact=True).count
            if actual != result["chunks"]:
                raise RuntimeError(f"Verification failed for {path}: expected {result['chunks']} chunks, found {actual}")

    for path, result in results.items():
        finish_pending_file(pending_path, "remaining_additions", path, result)

    summary = {
        "files_added": sum(result["status"] == "indexed" for result in results.values()),
        "files_skipped": sum(result["status"] == "skipped" for result in results.values()),
        "chunks_added": len(all_chunks),
        "input_tokens": usage["input_tokens"],
    }
    get_client().update_current_span(metadata=summary)
    return summary


@observe(name="ingestion", capture_input=False, capture_output=True)
def sync_changes(full_rebuild=False):
    """Resume pending work, or index all tracked files when full_rebuild is True."""
    state_dir = project_root / "data" / "ingestion_state"
    state_dir.mkdir(parents=True, exist_ok=True)
    pending_path = state_dir / "pending.json"
    checkpoint_path = state_dir / "checkpoint.json"

    run = {
        "run_id": str(uuid4()),
        "repository": str(repo),
        "collection": collection,
        "status": "started",
        "started_at": datetime.now(timezone.utc).isoformat(),
        "retrying_pending": pending_path.exists(),
    }
    run_path = state_dir / "runs" / f"{run['run_id']}.json"

    # The OS releases this lock even if the process crashes.
    with (state_dir / "sync.lock").open("a") as lock:
        try:
            fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError:
            raise RuntimeError("Another ingestion run is already active") from None

        save_json(run_path, run)
        get_client().update_current_span(metadata=run)

        try:
            checkpoint = (
                json.loads(checkpoint_path.read_text())
                if checkpoint_path.exists() else {}
            )
            report = sync_data(
                repo, checkpoint.get("indexed_commit"), pending_path,
                full_rebuild=full_rebuild,
            )
            run.update({
                "full_rebuild": report.get("full_rebuild", False),
                "old_commit": report["old_commit"],
                "target_commit": report["target_commit"],
                "file_counts": report["file_counts"],
            })

            paths_to_delete = [Path(p) for p in report["remaining_deletions"]]
            addition_candidates = [Path(p) for p in report["remaining_additions"]]
            paths_to_add = [
                path for path in filter_files(addition_candidates, repo)
                if not path.name.startswith("~$")
            ]
            accepted_paths = set(paths_to_add)
            skipped_paths = [path for path in addition_candidates if path not in accepted_paths]

            for path in skipped_paths:
                finish_pending_file(pending_path, "remaining_additions", path, {
                    "status": "skipped",
                    "reason": "temporary file or rejected by file filter",
                })
            run["files_skipped"] = len(skipped_paths)

            with get_client().start_as_current_observation(name="ingestion.delete_sources") as span:
                run["chunks_deleted"] = delete_sources(paths_to_delete, pending_path)
                span.update(output={"chunks_deleted": run["chunks_deleted"]})

            addition_summary = add_sources(paths_to_add, pending_path)
            addition_summary["files_skipped"] += len(skipped_paths)
            run.update(addition_summary)

            pending = json.loads(pending_path.read_text())
            if pending["remaining_deletions"] or pending["remaining_additions"]:
                raise RuntimeError("Pending files remain; checkpoint cannot advance")
            run["completed_files"] = pending["completed_files"]

            save_json(checkpoint_path, {
                "indexed_commit": report["target_commit"],
                "run_id": run["run_id"],
                "indexed_at": datetime.now(timezone.utc).isoformat(),
            })
            pending_path.unlink(missing_ok=True)
            run["status"] = "succeeded"

        except Exception as error:
            run.update(status="failed", error_type=type(error).__name__, error=str(error))
            raise
        finally:
            run["finished_at"] = datetime.now(timezone.utc).isoformat()
            save_json(run_path, run)
            get_client().update_current_span(metadata=run)

    return run


if __name__ == "__main__":
    try:
        print(json.dumps(sync_changes(), indent=2))
    finally:
        get_client().flush()
